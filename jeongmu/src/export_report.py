"""정무판단 결과를 공유용 보고서(Word·PDF)와 슬라이드(PPT)로 내보낸다. AI 호출 없음(0원).

- 입력: data/strategy/<run_id>.json (기본 latest.json), 여론 게이지(server.polls_board), 일정(data/calendar.json),
  사용자 검토 의견(data/strategy/reviews.json, 있으면 카드별 '유효/보조/제외'와 사유를 표시).
- 공유용이라 근거 번호([E12])와 원칙 코드(L22, W26 …)는 지우고, 출처는 기사 링크로만 남긴다.
  API 키·내부 경로·의뢰인 제약 파일 원문은 넣지 않는다.
- PDF는 Word 파일을 LibreOffice(soffice)로 변환한다. 없으면 Word 파일을 열어 PDF로 저장하면 된다.

사용: python src/export_report.py               # docx·pptx·pdf 를 data/reports/ 에 생성
      python src/export_report.py --run st_xxx  # 특정 실행 결과
"""
import argparse
import json
import re
import shutil
import subprocess
import sys
from datetime import date, timedelta
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))

from config import DATA

STRAT = DATA / "strategy"
OUT = DATA / "reports"
FONT = "Malgun Gothic"
NAVY, BRASS, GREY = (0x1B, 0x2A, 0x41), (0xB8, 0x86, 0x2E), (0x6B, 0x72, 0x80)
CODE = r"[DPLWFCST]\d{1,2}"


# ---------- 내용 정리 ----------

def clean(text: str) -> str:
    """근거 번호와 원칙 코드를 지운다. 코드가 남은 문장(내부 설명)은 문장째 뺀다."""
    t = re.sub(r"\[E\d+\]", "", text or "")
    t = re.sub(rf"\(\s*{CODE}[^)]*\)", "", t)
    parts = re.split(r"(?<=[.!?다])\s+", t)
    t = " ".join(p for p in parts if not re.search(rf"(?<![A-Za-z0-9]){CODE}(?![0-9])", p))
    return re.sub(r"\s{2,}", " ", t).replace(" .", ".").strip()


def load_run(run: str | None) -> dict:
    p = STRAT / (f"{run}.json" if run else "latest.json")
    if not p.exists():
        raise FileNotFoundError(f"전략 결과가 없습니다: {p.name}")
    return json.loads(p.read_text(encoding="utf-8"))


def load_reviews(run_id: str) -> dict:
    p = STRAT / "reviews.json"
    return json.loads(p.read_text(encoding="utf-8")).get(run_id, {}) if p.exists() else {}


def poll_latest() -> list[dict]:
    """조사기관별 최신 값(대통령 긍정·민주·국힘)."""
    try:
        import server
        b = server.polls_board()
    except Exception:
        return []
    rows: dict[str, dict] = {}
    for key, col in (("대통령국정:긍정", "pres"), ("정당:더불어민주당", "dem"), ("정당:국민의힘", "ppp")):
        for pl, xs in b.get(key, {}).items():
            r = rows.setdefault(pl, {"pollster": pl, "date": ""})
            r[col] = xs[-1]["value"]
            r["date"] = max(r["date"], xs[-1]["date"])
    return sorted(rows.values(), key=lambda r: r["pollster"])


def poll_series(pollster: str = "리얼미터") -> dict:
    try:
        import server
        b = server.polls_board()
    except Exception:
        return {}
    out = {}
    for key, name in (("정당:더불어민주당", "더불어민주당"), ("정당:국민의힘", "국민의힘"), ("대통령국정:긍정", "대통령 국정 긍정")):
        xs = b.get(key, {}).get(pollster)
        if xs:
            out[name] = {x["date"]: x["value"] for x in xs}
    return out


def upcoming(on: str, days: int = 45) -> list[dict]:
    p = DATA / "calendar.json"
    if not p.exists():
        return []
    hi = (date.fromisoformat(on) + timedelta(days=days)).isoformat()
    ev = [e for e in json.loads(p.read_text(encoding="utf-8"))["events"] if on <= e.get("end", e["date"]) and e["date"] <= hi]
    return sorted(ev, key=lambda e: e["date"])


def sources(d: dict, limit: int = 15) -> list[tuple[str, str]]:
    """보고서 본문이 기댄 근거 중 링크가 있는 것(여론조사·기사)."""
    used = set()
    for x in [d["situation"], *d["diagnosis"], *d["strategies"], *d["scenarios"]]:
        used.update(x.get("refs", []))
    out, seen = [], set()
    for e in d.get("evidence", []):
        if e["ref"] in used and e.get("link") and e["link"] not in seen:
            seen.add(e["link"])
            label = re.sub(r"^\[[^\]]+\]\s*", "", e["text"])[:70]
            out.append((label, e["link"]))
    return out[:limit]


def build(run: str | None = None) -> dict:
    d = load_run(run)
    rv = load_reviews(d.get("run_id", ""))
    cards = []
    for c in d["strategies"]:
        r = rv.get(c["name"], {})
        cards.append({**c, "verdict": r.get("verdict", "검토 전"), "review": r.get("note", ""),
                      "steps": [clean(s) for s in c["steps"]], "rationale": clean(c["rationale"]),
                      "opponent_counter": clean(c["opponent_counter"]),
                      "counter_to_counter": clean(c["counter_to_counter"]),
                      "principal_fit": clean(c.get("principal_fit", ""))})
    order = {"유효": 0, "보조": 1, "검토 전": 2, "제외": 3}
    cards.sort(key=lambda c: (order.get(c["verdict"], 2), 0 if c.get("scale") == "거시" else 1))
    return {
        "goal": d["goal"], "side": d["side"], "on": d["on"], "run_id": d.get("run_id", ""),
        "situation": clean(d["situation"]["text"]),
        "diagnosis": [{**x, "rationale": clean(x["rationale"])} for x in d["diagnosis"]],
        "cards": cards,
        "scenarios": sorted(({**x, "path": [clean(p) for p in x["path"]], "end_state": clean(x["end_state"])}
                             for x in d["scenarios"]), key=lambda x: -x["probability"]),
        "watch": [clean(w) for w in d["watch"]],
        "predictions": d["predictions"],
        "polls": poll_latest(), "series": poll_series(),
        "calendar": upcoming(d["on"]),
        "sources": sources(d),
    }


def headline(R: dict) -> list[str]:
    """한 장 요약 4줄."""
    sents = [x for x in re.split(r"(?<=다\.)\s+", R["situation"]) if not x.startswith("기준일")]
    lead = next((x for x in sents if "국면" in x), sents[0] if sents else "")
    lead = re.sub(r"^(그러나|하지만|다만)\s+", "", lead)
    top_down = next((x for x in R["diagnosis"] if x["direction"] == "하락 요인" and x["weight"] == "상"), None)
    keep = [c["name"] for c in R["cards"] if c["verdict"] in ("유효", "검토 전")][:3]
    sc = R["scenarios"][0] if R["scenarios"] else None
    out = [lead]
    if top_down:
        out.append(f"최대 감점 요인: {top_down['driver']}")
    if keep:
        out.append("핵심 카드: " + " / ".join(keep))
    if sc:
        out.append(f"가장 유력한 전개: {sc['name']}({round(sc['probability'] * 100)}%)")
    return out


def pct(v) -> str:
    return "—" if v is None else f"{v:g}%"


# ---------- Word ----------

def to_docx(R: dict, path: Path) -> Path:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name, st.font.size = FONT, Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    for s in ("Title", "Heading 1", "Heading 2"):
        doc.styles[s].font.name = FONT
        doc.styles[s].element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        doc.styles[s].font.color.rgb = RGBColor(*NAVY)

    def table(headers, rows, widths=None):
        t = doc.add_table(rows=1, cols=len(headers))
        t.style = "Light Grid Accent 1"
        for i, h in enumerate(headers):
            t.rows[0].cells[i].text = h
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = str(v)
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(9)
        doc.add_paragraph()
        return t

    doc.add_heading("GAMBIT 정무판단 보고서", 0)
    m = doc.add_paragraph(f"기준일 {R['on']} · 관점 {R['side']} · 실행 {R['run_id']}")
    m.runs[0].font.color.rgb = RGBColor(*GREY)
    doc.add_paragraph(f"과제: {R['goal']}")

    doc.add_heading("한 장 요약", 1)
    for line in headline(R):
        doc.add_paragraph(line, style="List Bullet")

    doc.add_heading("판세", 1)
    doc.add_paragraph(R["situation"])
    if R["polls"]:
        table(["조사기관", "대통령 국정 긍정", "민주당", "국민의힘", "최근 조사"],
              [[p["pollster"], pct(p.get("pres")), pct(p.get("dem")), pct(p.get("ppp")), p["date"]] for p in R["polls"]])
        doc.add_paragraph("조사기관마다 방식이 달라 같은 기관 안의 추세로만 비교합니다. 수치는 보도에서 자동 추출한 값입니다.").runs[0].font.size = Pt(9)

    doc.add_heading("진단", 1)
    table(["방향", "비중", "요인", "판단 근거"],
          [[x["direction"].replace(" 요인", ""), x["weight"], x["driver"], x["rationale"]] for x in R["diagnosis"]])

    doc.add_heading("전략 카드", 1)
    table(["카드", "구분", "기대 효과", "기간", "법적/역풍", "검토"],
          [[c["name"], f"{c.get('scale', '')} {c['kind']}".strip(), f"+{c['expected_pp_low']:g}~{c['expected_pp_high']:g}%p",
            f"{c['weeks']}주", f"{c['legal_risk']}/{c['backlash_risk']}",
            c["verdict"] + (f": {c['review']}" if c["review"] else "")] for c in R["cards"]])
    for c in R["cards"]:
        if c["verdict"] == "제외":
            continue
        doc.add_heading(c["name"], 2)
        doc.add_paragraph(f"{c.get('scale', '')} · {c['kind']} · 대상 {c['target_audience']} · 기대 +{c['expected_pp_low']:g}~{c['expected_pp_high']:g}%p / {c['weeks']}주")
        for s in c["steps"]:
            doc.add_paragraph(s, style="List Number")
        doc.add_paragraph(f"상대 대응: {c['opponent_counter']}")
        doc.add_paragraph(f"재반격: {c['counter_to_counter']}")
        if c["principal_fit"]:
            doc.add_paragraph(f"의뢰인 수용성: {c['principal_fit']}")
        if c["rationale"]:
            doc.add_paragraph(c["rationale"]).runs[0].font.color.rgb = RGBColor(*GREY)

    doc.add_heading("전개 시나리오", 1)
    table(["확률", "시나리오", "전개", "결말"],
          [[f"{round(x['probability'] * 100)}%", x["name"], " → ".join(x["path"]), x["end_state"]] for x in R["scenarios"]])

    if R["calendar"]:
        doc.add_heading("향후 일정", 1)
        table(["날짜", "일정", "상태"],
              [[e["date"] + (f"~{e['end']}" if e.get("end") else ""), e["title"], e["status"]] for e in R["calendar"]])

    doc.add_heading("주시 지표", 1)
    for w in R["watch"]:
        doc.add_paragraph(w, style="List Bullet")

    doc.add_heading("채점 예측", 1)
    table(["조사기관", "대상", "기한", "예측 범위"],
          [[p["pollster"], f"{p['target_kind']} {p['target']}", p["by_date"], f"{p['low']:g}~{p['high']:g}%"] for p in R["predictions"]])
    doc.add_paragraph("전략·확률·효과는 AI 추정입니다. 기한이 지나면 실제 조사로 채점합니다.").runs[0].font.size = Pt(9)

    if R["sources"]:
        doc.add_heading("출처", 1)
        for label, url in R["sources"]:
            doc.add_paragraph(f"{label} — {url}", style="List Bullet").runs[0].font.size = Pt(8.5)
    for p in doc.paragraphs:
        if p.style.name == "Title":
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path


# ---------- PPT ----------

def to_pptx(R: dict, path: Path) -> Path:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.util import Emu, Inches, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    W = prs.slide_width

    def txt(slide, x, y, w, h, text, size=14, bold=False, color=NAVY, lines=None):
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        for i, line in enumerate(lines if lines is not None else [text]):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            r = p.add_run()
            r.text = line
            r.font.size, r.font.bold, r.font.name = Pt(size), bold, FONT
            r.font.color.rgb = RGBColor(*color)
            p.space_after = Pt(6)
        return tb

    def page(title, kicker=""):
        s = prs.slides.add_slide(blank)
        bar = s.shapes.add_shape(1, 0, 0, W, Inches(0.12))
        bar.fill.solid()
        bar.fill.fore_color.rgb = RGBColor(*BRASS)
        bar.line.fill.background()
        if kicker:
            txt(s, 0.6, 0.35, 12, 0.4, kicker, 11, color=BRASS)
        txt(s, 0.6, 0.65, 12, 0.8, title, 26, bold=True)
        txt(s, 0.6, 7.0, 12, 0.3, f"GAMBIT · {R['on']} · AI 추정 포함", 9, color=GREY)
        return s

    def table(slide, x, y, w, headers, rows, size=11, col_w=None):
        shape = slide.shapes.add_table(len(rows) + 1, len(headers), Inches(x), Inches(y), Inches(w),
                                       Inches(0.4 * (len(rows) + 1)))
        t = shape.table
        if col_w:
            for i, cw in enumerate(col_w):
                t.columns[i].width = Inches(cw)
        for j, h in enumerate(headers):
            t.cell(0, j).text = h
        for i, row in enumerate(rows, 1):
            for j, v in enumerate(row):
                t.cell(i, j).text = str(v)
        for i in range(len(rows) + 1):
            for j in range(len(headers)):
                for p in t.cell(i, j).text_frame.paragraphs:
                    for r in p.runs:
                        r.font.size, r.font.name = Pt(size), FONT
        return t

    # 1 표지
    s = prs.slides.add_slide(blank)
    bg = s.shapes.add_shape(1, 0, 0, W, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(*NAVY)
    bg.line.fill.background()
    txt(s, 0.8, 2.2, 11, 1, "GAMBIT", 54, True, BRASS)
    txt(s, 0.8, 3.3, 11, 0.8, "정무판단 보고서", 30, True, (0xFF, 0xFF, 0xFF))
    txt(s, 0.8, 4.3, 11.5, 1, R["goal"], 16, color=(0xD8, 0xDE, 0xE6))
    txt(s, 0.8, 5.3, 11, 0.5, f"기준일 {R['on']} · 관점 {R['side']}", 12, color=(0xB0, 0xB8, 0xC4))

    # 2 한 장 요약
    s = page("한 장 요약", "BOTTOM LINE")
    txt(s, 0.6, 1.7, 12, 4.8, "", 18, lines=["• " + x for x in headline(R)])

    # 3 판세
    s = page("판세: 조사기관별 지지율", "POLLS")
    ser = R["series"]
    if ser:
        dates = sorted({d for v in ser.values() for d in v})
        cd = CategoryChartData()
        cd.categories = [d[5:].replace("-", "/") for d in dates]
        for name, vals in ser.items():
            cd.add_series(name, [vals.get(d) for d in dates])
        ch = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(0.6), Inches(1.6), Inches(7.4), Inches(5.1), cd).chart
        ch.has_legend = True
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM
        ch.legend.include_in_layout = False
        ch.legend.font.size, ch.legend.font.name = Pt(11), FONT
        ch.category_axis.tick_labels.font.size = Pt(9)
        ch.value_axis.tick_labels.font.size = Pt(9)
        vals = [v for x in ser.values() for v in x.values()]
        ch.value_axis.minimum_scale = max(0, (int(min(vals)) // 5) * 5 - 5)
        ch.value_axis.maximum_scale = (int(max(vals)) // 5) * 5 + 5
        colors = {"더불어민주당": (0x1F, 0x5F, 0xBF), "국민의힘": (0xD6, 0x33, 0x3A), "대통령 국정 긍정": BRASS}
        for plot_ser in ch.plots[0].series:
            c = colors.get(plot_ser.name)
            if c:
                plot_ser.format.line.color.rgb = RGBColor(*c)
                plot_ser.marker.format.fill.solid()
                plot_ser.marker.format.fill.fore_color.rgb = RGBColor(*c)
        txt(s, 0.6, 1.3, 7, 0.3, "리얼미터 추이(%)", 11, color=GREY)
    if R["polls"]:
        table(s, 8.3, 1.7, 4.5, ["기관", "대통령", "민주", "국힘"],
              [[p["pollster"].replace("(전국지표조사)", ""), pct(p.get("pres")), pct(p.get("dem")), pct(p.get("ppp"))] for p in R["polls"]],
              12, [1.5, 1.0, 1.0, 1.0])
        txt(s, 8.3, 3.6, 4.5, 1.2, "", 10, color=GREY,
            lines=["기관마다 조사 방식이 달라 같은 기관 추세로만 비교", "보도에서 자동 추출한 수치"])

    # 4 진단
    s = page("진단: 무엇이 지지율을 움직이나", "DRIVERS")
    pips = {"상": "●●●", "중": "●●○", "하": "●○○"}
    down = [f"{pips.get(x['weight'], '')}  {x['driver']}" for x in R["diagnosis"] if x["direction"] == "하락 요인"]
    up = [f"{pips.get(x['weight'], '')}  {x['driver']}" for x in R["diagnosis"] if x["direction"] == "상승 요인"]
    txt(s, 0.6, 1.6, 6.2, 0.4, "하락 요인", 16, True, (0xD6, 0x33, 0x3A))
    txt(s, 0.6, 2.1, 6.2, 4.6, "", 13, lines=down)
    txt(s, 7.0, 1.6, 5.8, 0.4, "상승 요인", 16, True, (0x1E, 0x8E, 0x5A))
    txt(s, 7.0, 2.1, 5.8, 4.6, "", 13, lines=up)

    # 5 카드 요약
    s = page("전략 카드 한눈에", "PLAYS")
    table(s, 0.6, 1.6, 12.1, ["카드", "구분", "기대 효과", "기간", "법/역풍", "검토"],
          [[c["name"], f"{c.get('scale', '')} {c['kind']}".strip(), f"+{c['expected_pp_low']:g}~{c['expected_pp_high']:g}%p",
            f"{c['weeks']}주", f"{c['legal_risk']}/{c['backlash_risk']}", c["verdict"]] for c in R["cards"]],
          11, [4.6, 1.7, 1.6, 0.8, 1.0, 2.4])

    # 6.. 카드별
    for c in R["cards"]:
        if c["verdict"] == "제외":
            continue
        s = page(c["name"], f"{c.get('scale', '')} · {c['kind']}  |  +{c['expected_pp_low']:g}~{c['expected_pp_high']:g}%p · {c['weeks']}주 · 법 {c['legal_risk']} / 역풍 {c['backlash_risk']}")
        txt(s, 0.6, 1.6, 7.2, 5.2, "", 12, lines=[f"{i}. {x}" for i, x in enumerate(c["steps"][:5], 1)])
        txt(s, 8.1, 1.6, 4.7, 0.3, "상대 응수", 12, True, (0xD6, 0x33, 0x3A))
        txt(s, 8.1, 1.95, 4.7, 2.0, c["opponent_counter"][:220], 11)
        txt(s, 8.1, 4.1, 4.7, 0.3, "재반격", 12, True, (0x1E, 0x8E, 0x5A))
        txt(s, 8.1, 4.45, 4.7, 2.3, c["counter_to_counter"][:240], 11)

    # 시나리오
    s = page("전개 시나리오", "SCENARIOS")
    cd = CategoryChartData()
    sc = list(reversed(R["scenarios"]))
    cd.categories = [x["name"] for x in sc]
    cd.add_series("확률", [round(x["probability"] * 100) for x in sc])
    ch = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.6), Inches(1.6), Inches(5.6), Inches(4.8), cd).chart
    ch.has_legend = False
    ch.has_title = False
    ch.plots[0].has_data_labels = True
    ch.plots[0].data_labels.font.size = Pt(12)
    ch.plots[0].data_labels.number_format = '0"%"'
    ch.plots[0].data_labels.number_format_is_linked = False
    ch.category_axis.tick_labels.font.size = Pt(11)
    ch.value_axis.visible = False
    ch.plots[0].series[0].format.fill.solid()
    ch.plots[0].series[0].format.fill.fore_color.rgb = RGBColor(*NAVY)
    txt(s, 6.6, 1.6, 6.2, 5.2, "", 11,
        lines=[f"{x['name']}: {x['end_state'][:110]}" for x in R["scenarios"]])

    # 일정
    if R["calendar"]:
        s = page("향후 일정", "CALENDAR")
        rows = [[e["date"][5:].replace("-", "/") + (f"~{e['end'][5:].replace('-', '/')}" if e.get("end") else ""),
                 e["title"][:60], e["status"]] for e in R["calendar"][:12]]
        table(s, 0.6, 1.6, 12.1, ["날짜", "일정", "상태"], rows, 11, [1.8, 8.9, 1.4])

    # 예측
    s = page("채점 예측과 주시 지표", "SCORING")
    table(s, 0.6, 1.6, 6.4, ["기관", "대상", "11/30 범위"],
          [[p["pollster"].replace("(전국지표조사)", ""), f"{p['target']}", f"{p['low']:g}~{p['high']:g}%"] for p in R["predictions"]],
          11, [1.6, 3.0, 1.8])
    txt(s, 7.3, 1.6, 5.5, 5.2, "", 11, lines=["• " + w[:70] for w in R["watch"][:8]])

    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(path)
    return path


def to_pdf(docx_path: Path) -> Path | None:
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        return None
    subprocess.run([exe, "--headless", "--convert-to", "pdf", "--outdir", str(docx_path.parent), str(docx_path)],
                   check=True, capture_output=True, timeout=180)
    pdf = docx_path.with_suffix(".pdf")
    return pdf if pdf.exists() else None


def export(fmt: str, run: str | None = None) -> Path:
    R = build(run)
    stem = OUT / f"GAMBIT_{R['on']}_{R['run_id'] or 'latest'}"
    if fmt == "docx":
        return to_docx(R, stem.with_suffix(".docx"))
    if fmt == "pptx":
        return to_pptx(R, stem.with_suffix(".pptx"))
    if fmt == "pdf":
        pdf = to_pdf(to_docx(R, stem.with_suffix(".docx")))
        if not pdf:
            raise RuntimeError("PDF 변환에 LibreOffice가 필요합니다. Word 파일을 열어 PDF로 저장하세요.")
        return pdf
    raise ValueError("fmt 는 docx, pptx, pdf 중 하나")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--run")
    ap.add_argument("--fmt", nargs="*", default=["docx", "pptx", "pdf"])
    a = ap.parse_args()
    for f in a.fmt:
        print(export(f, a.run))
